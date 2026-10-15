import io
import openpyxl
from app.export.excel_generator import excel_generator

def test_generate_5_sheet_excel():
    records = [
        {
            "name": "K. RAHUL REDDY",
            "roll_number": "0020671A0512", # Test leading zeros!
            "college_name": "JB Institute of Engineering",
            "course": "B.Tech",
            "branch": "CSE",
            "batch": "2020-2024",
            "latest_cgpa": 8.65,
            "semester_results": [
                {
                    "semester": "III-II",
                    "regulation": "R18",
                    "examination_name": "B.Tech Regular Exam",
                    "sgpa": 8.58,
                    "cgpa": 8.65,
                    "backlog_count": 0,
                    "overall_result": "PASSED",
                    "verification_status": "VERIFIED",
                    "source_filename": "memo_1.pdf",
                    "subject_marks": [
                        {
                            "subject_code": "18CS601",
                            "subject_name": "=MALICIOUS_CMD()", # Formula injection attempt
                            "internal_marks": 26,
                            "external_marks": 48,
                            "total_marks": 74,
                            "max_marks": 100,
                            "grade": "A",
                            "grade_points": 8.0,
                            "credits": 3.0,
                            "subject_result": "P",
                            "is_pass": True
                        }
                    ],
                    "issues": []
                }
            ]
        }
    ]

    excel_bytes = excel_generator.generate_workbook(records)
    assert len(excel_bytes) > 0

    # Load with openpyxl to verify valid workbook
    wb = openpyxl.load_workbook(io.BytesIO(excel_bytes))
    
    # Assert all 5 sheets exist
    expected_sheets = ["Student Details", "Subject Marks", "Backlogs", "Summary", "Processing Issues"]
    for s_name in expected_sheets:
        assert s_name in wb.sheetnames

    # Verify formula injection escaping
    ws_sub = wb["Subject Marks"]
    escaped_sub_name = ws_sub.cell(row=2, column=5).value
    assert escaped_sub_name.startswith("'=") # Prepended apostrophe!

    # Verify roll number format preserved as text with leading zeros
    ws_stud = wb["Student Details"]
    roll_val = ws_stud.cell(row=2, column=2).value
    assert str(roll_val) == "0020671A0512"
