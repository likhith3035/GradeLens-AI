import httpx
from pathlib import Path
import openpyxl

def run_e2e_verification():
    base = "http://127.0.0.1:8000"
    sample_dir = Path(__file__).resolve().parent.parent.parent / "sample_memos"
    client = httpx.Client(timeout=30.0)

    print("--- STEP 1: Uploading 4 synthetic marks memos ---")
    files = [
        ("files", ("memo_1_jntu_passed.pdf", open(sample_dir / "memo_1_jntu_passed.pdf", "rb"), "application/pdf")),
        ("files", ("memo_2_jntu_backlog.pdf", open(sample_dir / "memo_2_jntu_backlog.pdf", "rb"), "application/pdf")),
        ("files", ("memo_3_anna_univ.pdf", open(sample_dir / "memo_3_anna_univ.pdf", "rb"), "application/pdf")),
        ("files", ("memo_4_flagged_review.pdf", open(sample_dir / "memo_4_flagged_review.pdf", "rb"), "application/pdf")),
    ]

    res = client.post(f"{base}/api/upload", files=files, data={"profile": "auto"})
    assert res.status_code == 200, res.text
    batch = res.json()
    print(f"Batch Upload Result: Total={batch['total_files']}, Processed={batch['processed_files']}, NeedsReview={batch['needs_review_files']}")

    print("\n--- STEP 2: Fetching Student Directory ---")
    s_res = client.get(f"{base}/api/students")
    assert s_res.status_code == 200
    students = s_res.json()
    print(f"Total Unique Students Indexed: {len(students)}")
    for s in students:
        print(f" - Roll: {s['roll_number']} | Name: {s['name']} | Branch: {s['branch']} | Backlogs: {s['total_backlogs']} | CGPA: {s['latest_cgpa']}")

    print("\n--- STEP 3: Checking Verification Queue ---")
    q_res = client.get(f"{base}/api/review/queue")
    queue = q_res.json()
    print(f"Items in Review Queue: {len(queue)}")
    for item in queue:
        first_issue = item['issues'][0]['message'] if item['issues'] else "N/A"
        print(f" - Flagged: Roll={item['roll_number']}, Issues={item['issues_count']}, Reason={first_issue}")

    if queue:
        target_id = queue[0]["result_id"]
        print(f"\n--- STEP 4: Simulating Manual Review Correction on ID {target_id} ---")
        corr_res = client.post(f"{base}/api/review/{target_id}/correct", json={
            "student_name": "V. SNEHA REDDY (VERIFIED)",
            "roll_number": "20B91A0499",
            "branch": "ELECTRONICS & COMMUNICATION",
            "semester": "IV-I",
            "sgpa": 7.62,
            "verification_status": "VERIFIED",
            "actor": "Academic Administrator"
        })
        assert corr_res.status_code == 200
        print("Correction applied:", corr_res.json())

    print("\n--- STEP 5: Verifying Analytics KPI Aggregation ---")
    a_res = client.get(f"{base}/api/analytics")
    assert a_res.status_code == 200
    analytics = a_res.json()
    print("Overview Stats:", analytics["overview"])
    print("Grade Frequency:", analytics["grade_distribution"])

    print("\n--- STEP 6: Downloading and Validating 5-Sheet Excel Workbook ---")
    exp_res = client.get(f"{base}/api/export/excel")
    assert exp_res.status_code == 200
    assert exp_res.headers["content-type"] == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    
    excel_path = Path(__file__).resolve().parent.parent / "data" / "exports" / "test_report.xlsx"
    excel_path.parent.mkdir(parents=True, exist_ok=True)
    excel_path.write_bytes(exp_res.content)
    print(f"Workbook saved ({excel_path.stat().st_size} bytes)")

    wb = openpyxl.load_workbook(excel_path)
    print("Workbook Sheets:", wb.sheetnames)
    assert len(wb.sheetnames) == 5
    for sheet_name in ["Student Details", "Subject Marks", "Backlogs", "Summary", "Processing Issues"]:
        assert sheet_name in wb.sheetnames, f"Missing sheet: {sheet_name}"

    print("\n>>> ALL END-TO-END PIPELINE CHECKS COMPLETED SUCCESSFULLY! <<<")

if __name__ == "__main__":
    run_e2e_verification()
