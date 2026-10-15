import io
from typing import List, Dict, Any, Optional
from datetime import datetime
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from ..utils.security import escape_excel_formula

class ExcelExportGenerator:
    def __init__(self):
        # Professional styling definitions
        self.header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
        self.header_fill = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid")  # Navy / Indigo
        self.header_alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

        self.sub_header_fill = PatternFill(start_color="2563EB", end_color="2563EB", fill_type="solid")
        
        self.title_font = Font(name="Calibri", size=14, bold=True, color="1E3A8A")
        
        self.thin_border = Border(
            left=Side(style='thin', color='D1D5DB'),
            right=Side(style='thin', color='D1D5DB'),
            top=Side(style='thin', color='D1D5DB'),
            bottom=Side(style='thin', color='D1D5DB')
        )
        
        self.zebra_fill = PatternFill(start_color="F9FAFB", end_color="F9FAFB", fill_type="solid")
        self.alert_fill = PatternFill(start_color="FEF2F2", end_color="FEF2F2", fill_type="solid")
        self.success_fill = PatternFill(start_color="F0FDF4", end_color="F0FDF4", fill_type="solid")

    def generate_workbook(
        self,
        student_records: List[Dict[str, Any]],
        filter_metadata: Optional[Dict[str, Any]] = None
    ) -> bytes:
        """
        Builds a comprehensive 5-sheet consolidated academic Excel workbook.
        """
        wb = openpyxl.Workbook()
        
        # 1. Sheet 1: Student Details
        ws_students = wb.active
        ws_students.title = "Student Details"
        self._build_student_details_sheet(ws_students, student_records)

        # 2. Sheet 2: Subject Marks
        ws_subjects = wb.create_sheet(title="Subject Marks")
        self._build_subject_marks_sheet(ws_subjects, student_records)

        # 3. Sheet 3: Backlogs
        ws_backlogs = wb.create_sheet(title="Backlogs")
        self._build_backlogs_sheet(ws_backlogs, student_records)

        # 4. Sheet 4: Summary
        ws_summary = wb.create_sheet(title="Summary")
        self._build_summary_sheet(ws_summary, student_records, filter_metadata)

        # 5. Sheet 5: Processing Issues
        ws_issues = wb.create_sheet(title="Processing Issues")
        self._build_issues_sheet(ws_issues, student_records)

        output = io.BytesIO()
        wb.save(output)
        output.seek(0)
        return output.getvalue()

    def _style_headers(self, ws, row_idx: int, num_cols: int):
        ws.row_dimensions[row_idx].height = 28
        for col in range(1, num_cols + 1):
            cell = ws.cell(row=row_idx, column=col)
            cell.font = self.header_font
            cell.fill = self.header_fill
            cell.alignment = self.header_alignment
            cell.border = self.thin_border

    def _auto_fit_columns(self, ws, max_len_cap: int = 40):
        for col in ws.columns:
            max_len = 0
            col_letter = get_column_letter(col[0].column)
            for cell in col:
                val = str(cell.value or "")
                if len(val) > max_len and not cell.coordinate in ws.merged_cells:
                    max_len = len(val)
            adjusted_width = max(max_len + 3, 12)
            ws.column_dimensions[col_letter].width = min(adjusted_width, max_len_cap)

    def _build_student_details_sheet(self, ws, student_records: List[Dict[str, Any]]):
        headers = [
            "Student Name", "Roll Number", "College", "Course", "Branch", "Batch",
            "Examination", "Semester", "SGPA", "CGPA", "Current Backlogs",
            "Result", "Source Filename", "Verification Status"
        ]
        ws.append(headers)
        self._style_headers(ws, 1, len(headers))
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}1"

        row_num = 2
        for s in student_records:
            roll = str(s.get("roll_number", ""))
            sem_list = s.get("semester_results", [])
            if not sem_list:
                sem_list = [{}]

            for sem in sem_list:
                row_data = [
                    escape_excel_formula(s.get("name")),
                    escape_excel_formula(roll),
                    escape_excel_formula(s.get("college_name")),
                    escape_excel_formula(s.get("course")),
                    escape_excel_formula(s.get("branch")),
                    escape_excel_formula(s.get("batch")),
                    escape_excel_formula(sem.get("examination_name")),
                    escape_excel_formula(sem.get("semester")),
                    sem.get("sgpa") if sem.get("sgpa") is not None else "",
                    sem.get("cgpa") if sem.get("cgpa") is not None else (s.get("latest_cgpa") or ""),
                    sem.get("backlog_count", 0),
                    escape_excel_formula(sem.get("overall_result")),
                    escape_excel_formula(sem.get("source_filename", s.get("source_filename", ""))),
                    escape_excel_formula(sem.get("verification_status", s.get("verification_status", "VERIFIED")))
                ]
                ws.append(row_data)

                # Format roll number cell explicitly as text
                roll_cell = ws.cell(row=row_num, column=2)
                roll_cell.number_format = '@'

                # Apply styling & borders
                for c_idx in range(1, len(headers) + 1):
                    c = ws.cell(row=row_num, column=c_idx)
                    c.border = self.thin_border
                    if row_num % 2 == 1:
                        c.fill = self.zebra_fill
                
                row_num += 1

        self._auto_fit_columns(ws)

    def _build_subject_marks_sheet(self, ws, student_records: List[Dict[str, Any]]):
        headers = [
            "Roll Number", "Student Name", "Semester", "Subject Code", "Subject Name",
            "Internal Marks", "External Marks", "Practical Marks", "Total Marks",
            "Maximum Marks", "Grade", "Grade Points", "Credits", "Result"
        ]
        ws.append(headers)
        self._style_headers(ws, 1, len(headers))
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}1"

        row_num = 2
        for s in student_records:
            roll = str(s.get("roll_number", ""))
            sname = s.get("name", "")

            for sem in s.get("semester_results", []):
                sem_name = sem.get("semester", "")
                for sub in sem.get("subject_marks", []):
                    row_data = [
                        escape_excel_formula(roll),
                        escape_excel_formula(sname),
                        escape_excel_formula(sem_name),
                        escape_excel_formula(sub.get("subject_code")),
                        escape_excel_formula(sub.get("subject_name")),
                        sub.get("internal_marks") if sub.get("internal_marks") is not None else "",
                        sub.get("external_marks") if sub.get("external_marks") is not None else "",
                        sub.get("practical_marks") if sub.get("practical_marks") is not None else "",
                        sub.get("total_marks") if sub.get("total_marks") is not None else "",
                        sub.get("max_marks") if sub.get("max_marks") is not None else "",
                        escape_excel_formula(sub.get("grade")),
                        sub.get("grade_points") if sub.get("grade_points") is not None else "",
                        sub.get("credits") if sub.get("credits") is not None else "",
                        escape_excel_formula(sub.get("subject_result"))
                    ]
                    ws.append(row_data)

                    # Text format for roll & subject code
                    ws.cell(row=row_num, column=1).number_format = '@'
                    ws.cell(row=row_num, column=4).number_format = '@'

                    for c_idx in range(1, len(headers) + 1):
                        c = ws.cell(row=row_num, column=c_idx)
                        c.border = self.thin_border
                        if sub.get("is_pass") is False:
                            c.fill = self.alert_fill
                        elif row_num % 2 == 1:
                            c.fill = self.zebra_fill

                    row_num += 1

        self._auto_fit_columns(ws)

    def _build_backlogs_sheet(self, ws, student_records: List[Dict[str, Any]]):
        headers = [
            "Roll Number", "Student Name", "Branch", "Semester", "Subject Code",
            "Subject Name", "Grade", "Status", "Backlog Type"
        ]
        ws.append(headers)
        self._style_headers(ws, 1, len(headers))
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}1"

        row_num = 2
        for s in student_records:
            roll = str(s.get("roll_number", ""))
            sname = s.get("name", "")
            branch = s.get("branch", "")

            for sem in s.get("semester_results", []):
                sem_name = sem.get("semester", "")
                for sub in sem.get("subject_marks", []):
                    if sub.get("is_pass") is False:
                        row_data = [
                            escape_excel_formula(roll),
                            escape_excel_formula(sname),
                            escape_excel_formula(branch),
                            escape_excel_formula(sem_name),
                            escape_excel_formula(sub.get("subject_code")),
                            escape_excel_formula(sub.get("subject_name")),
                            escape_excel_formula(sub.get("grade")),
                            escape_excel_formula(sub.get("subject_result", "FAIL")),
                            "Current Semester Failure"
                        ]
                        ws.append(row_data)

                        ws.cell(row=row_num, column=1).number_format = '@'
                        ws.cell(row=row_num, column=5).number_format = '@'

                        for c_idx in range(1, len(headers) + 1):
                            c = ws.cell(row=row_num, column=c_idx)
                            c.border = self.thin_border
                            c.fill = self.alert_fill

                        row_num += 1

        self._auto_fit_columns(ws)

    def _build_summary_sheet(self, ws, student_records: List[Dict[str, Any]], filter_metadata: Optional[Dict[str, Any]]):
        total_students = len(student_records)
        total_memos = sum(len(s.get("semester_results", [])) for s in student_records)
        all_subjects = [sub for s in student_records for sem in s.get("semester_results", []) for sub in sem.get("subject_marks", [])]
        total_subjects = len(all_subjects)
        total_failed_subs = sum(1 for sub in all_subjects if sub.get("is_pass") is False)
        pass_pct = round(((total_subjects - total_failed_subs) / total_subjects * 100), 2) if total_subjects > 0 else 0.0

        valid_cgpas = [s.get("latest_cgpa") for s in student_records if s.get("latest_cgpa") is not None]
        avg_cgpa = round(sum(valid_cgpas) / len(valid_cgpas), 2) if valid_cgpas else "N/A"

        # Title
        ws.cell(row=1, column=1, value="MEMOSCAN AI — ACADEMIC PERFORMANCE CONSOLIDATED REPORT").font = self.title_font
        ws.cell(row=2, column=1, value=f"Generated on: {datetime.utcnow().strftime('%Y-%m-%d %H:%M:%S')} UTC | Total Memos: {total_memos}")
        
        # Summary KPI Box
        summary_kpis = [
            ("Total Unique Students", total_students),
            ("Total Semester Memos Processed", total_memos),
            ("Total Subject Records", total_subjects),
            ("Total Passed Subject Records", total_subjects - total_failed_subs),
            ("Total Failed Records (Backlogs)", total_failed_subs),
            ("Overall Subject Pass Percentage", f"{pass_pct}%"),
            ("Average Student CGPA", avg_cgpa),
        ]

        ws.cell(row=4, column=1, value="Metric").font = self.header_font
        ws.cell(row=4, column=1).fill = self.sub_header_fill
        ws.cell(row=4, column=2, value="Value").font = self.header_font
        ws.cell(row=4, column=2).fill = self.sub_header_fill

        for idx, (metric, val) in enumerate(summary_kpis, start=5):
            ws.cell(row=idx, column=1, value=metric).border = self.thin_border
            ws.cell(row=idx, column=2, value=str(val)).border = self.thin_border

        # Branch breakdown table
        start_row = 5 + len(summary_kpis) + 2
        ws.cell(row=start_row, column=1, value="Branch Breakdown").font = Font(size=12, bold=True, color="1E3A8A")
        
        b_headers = ["Branch", "Students", "Total Memos", "Passed Memos", "Pass %", "Avg CGPA"]
        ws.row_dimensions[start_row + 1].height = 22
        for col_i, h in enumerate(b_headers, start=1):
            c = ws.cell(row=start_row + 1, column=col_i, value=h)
            c.font = self.header_font
            c.fill = self.header_fill
            c.border = self.thin_border

        # Aggregate by branch
        branch_map = {}
        for s in student_records:
            b = s.get("branch") or "General / Unspecified"
            if b not in branch_map:
                branch_map[b] = {"students": 0, "memos": 0, "passed": 0, "cgpas": []}
            branch_map[b]["students"] += 1
            for sem in s.get("semester_results", []):
                branch_map[b]["memos"] += 1
                if sem.get("backlog_count", 0) == 0:
                    branch_map[b]["passed"] += 1
            if s.get("latest_cgpa") is not None:
                branch_map[b]["cgpas"].append(s.get("latest_cgpa"))

        cur_row = start_row + 2
        for b_name, b_data in branch_map.items():
            b_memos = b_data["memos"]
            b_pct = round((b_data["passed"] / b_memos * 100), 1) if b_memos > 0 else 0.0
            b_avg = round(sum(b_data["cgpas"]) / len(b_data["cgpas"]), 2) if b_data["cgpas"] else "-"

            ws.append([
                escape_excel_formula(b_name),
                b_data["students"],
                b_memos,
                b_data["passed"],
                f"{b_pct}%",
                b_avg
            ])
            for ci in range(1, len(b_headers) + 1):
                ws.cell(row=cur_row, column=ci).border = self.thin_border
            cur_row += 1

        self._auto_fit_columns(ws)

    def _build_issues_sheet(self, ws, student_records: List[Dict[str, Any]]):
        headers = [
            "Source Filename", "Roll Number", "Student Name", "Severity",
            "Issue Type", "Field Affected", "Description / Message", "Status"
        ]
        ws.append(headers)
        self._style_headers(ws, 1, len(headers))
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}1"

        row_num = 2
        for s in student_records:
            roll = str(s.get("roll_number", ""))
            sname = s.get("name", "")

            for sem in s.get("semester_results", []):
                fname = sem.get("source_filename", s.get("source_filename", ""))
                for issue in sem.get("issues", []):
                    row_data = [
                        escape_excel_formula(fname),
                        escape_excel_formula(roll),
                        escape_excel_formula(sname),
                        escape_excel_formula(issue.get("severity", "WARNING")),
                        escape_excel_formula(issue.get("issue_type")),
                        escape_excel_formula(issue.get("field_name")),
                        escape_excel_formula(issue.get("message")),
                        "Resolved" if issue.get("is_resolved") else "Needs Review"
                    ]
                    ws.append(row_data)

                    ws.cell(row=row_num, column=2).number_format = '@'

                    for c_idx in range(1, len(headers) + 1):
                        c = ws.cell(row=row_num, column=c_idx)
                        c.border = self.thin_border
                        if issue.get("severity") == "ERROR":
                            c.fill = self.alert_fill

                    row_num += 1

        self._auto_fit_columns(ws)

excel_generator = ExcelExportGenerator()
